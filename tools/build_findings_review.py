"""Build the findings review workbook (Engineer -> Architect -> Zero Trust Architect -> Client).

Usage:
    python tools/build_findings_review.py --register docs/findings-register.md --out workbooks/findings-review.xlsx
    python tools/build_findings_review.py --register docs/findings-register.md --gaps docs/gap-register.md --out workbooks/findings-review.xlsx
    python tools/build_findings_review.py --out findings-review.template.xlsx          # blank, client-agnostic template

The register is the Markdown table in findings-register.md. Only the Engineer columns are pre-filled;
every other section is left for its owner. The optional gap register (rows starting "| G-nn |", columns in
GAP_COLS order) fills the Gaps sheet: control, capability, evidence, coverage and decision gaps in one place.
"""
import argparse
import re
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

SECTIONS = [
    # (section title, owner, header fill, light fill, columns[(name, width, comment)])
    ('Engineer finding', 'Identity Engineer', '1F4E79', 'DDEBF7', [
        ('ID', 7, 'Finding ID, F-01 onward'),
        ('Finding', 34, 'What was observed, in one line'),
        ('Source and filter', 34, 'Script, file and filter or summary metric that produced the count'),
        ('Count', 16, 'Observed count'),
        ('Population', 22, 'What the count is out of (the denominator)'),
        ('Observed severity', 12, 'Critical / High / Medium / Low, per the Definitions sheet'),
        ('HITRUST CSF', 22, 'HITRUST CSF control reference'),
        ('Proposed ZT pillar / principle', 26, "Engineer's proposal; the Zero Trust Architect confirms"),
        ('Initial recommendation', 36, "Engineer's first recommendation; the Architect owns the solution"),
        ('Evidence note', 18, 'Evidence note in the evidence location (findings\\F-xx.md): run folder, SHA-256, run account, closure test'),
        ('Engineer date', 12, 'Date the finding was drafted'),
    ]),
    ('Architect assessment and solution', 'Identity Architect', '375623', 'E2EFDA', [
        ('Architect severity', 12, 'Agree or adjust the observed severity'),
        ('Root cause', 30, 'Why this exists (process, design, tooling, ownership)'),
        ('Proposed solution', 40, 'What to change, at the level a change record can be built from'),
        ('Effort', 9, 'S / M / L / XL'),
        ('Dependencies', 24, 'Access, decisions, other findings or platforms this depends on'),
        ('Target gate / sprint', 14, 'SOW gate (G1-G5) or sprint'),
        ('Architect', 16, 'Reviewer (role or name)'),
        ('Architect date', 12, 'Date reviewed'),
    ]),
    ('Zero Trust alignment', 'Zero Trust Architect', '7030A0', 'E4DFEC', [
        ('ZT pillar', 18, 'CISA Zero Trust Maturity Model pillar or cross-cutting capability'),
        ('ZT principle', 20, 'Verify explicitly / Use least privilege / Assume breach'),
        ('Current maturity', 12, 'Traditional / Initial / Advanced / Optimal'),
        ('Target maturity', 12, 'Maturity the solution reaches'),
        ('Alignment', 16, 'Does the proposed solution fit the target Zero Trust design?'),
        ('ZT comment', 30, 'Changes needed for alignment, or why not aligned'),
        ('ZT Architect', 16, 'Reviewer (role or name)'),
        ('ZT date', 12, 'Date reviewed'),
    ]),
    ('Client response', 'Client decision authority', 'C55A11', 'FCE4D6', [
        ('Client decision', 18, 'Accept and remediate / Accept risk / Defer / Reject'),
        ('Decided by', 16, 'Client decision authority (role or name)'),
        ('Decision date', 12, 'Date of the decision'),
        ('Client comment', 32, 'Conditions, rationale, or reason for risk acceptance, deferral or rejection'),
    ]),
    ('Disposition', 'Calculated / tracking', '404040', 'EDEDED', [
        ('Stage', 16, 'Calculated from which sections are complete'),
        ('Waiting on', 18, 'Calculated: who acts next'),
        ('Disposition', 30, 'Calculated from the client decision: where the finding goes next'),
        ('Tracking ID', 16, 'Jira story, exception register ID or status register ID'),
    ]),
]

LISTS = {
    'Severity': ['Critical', 'High', 'Medium', 'Low'],
    'Effort': ['S', 'M', 'L', 'XL'],
    'Gate': ['G0 Pre-project', 'G1', 'G2', 'G3', 'G4', 'G5', 'Follow-on'],
    'Pillar': ['Identity', 'Devices', 'Networks', 'Applications and workloads', 'Data',
               'Visibility and analytics', 'Automation and orchestration', 'Governance'],
    'Principle': ['Verify explicitly', 'Use least privilege', 'Assume breach'],
    'Maturity': ['Traditional', 'Initial', 'Advanced', 'Optimal'],
    'Alignment': ['Aligned', 'Aligned with changes', 'Not aligned'],
    'Decision': ['Accept and remediate', 'Accept risk', 'Defer', 'Reject'],
    'GapType': ['Control', 'Capability', 'Evidence', 'Coverage', 'Decision'],
    'GapStatus': ['Open', 'In progress', 'Closed', 'Accepted'],
    'Priority': ['High', 'Medium', 'Low'],
}

GAP_COLS = [
    # (name, width, comment, list)
    ('ID', 7, 'Gap ID, G-01 onward', None),
    ('Type', 12, 'Control = environment falls short (theme over findings); Capability = process or tool missing; Evidence = could not be assessed yet; Coverage = nothing measures it; Decision = waiting on a decision', 'GapType'),
    ('Gap', 34, 'The gap, in one line', None),
    ('Area', 16, 'Zero Trust pillar or domain', 'Pillar'),
    ('Current state', 30, 'What is in place today', None),
    ('Target state', 30, 'What closes the gap', None),
    ('Impact', 30, 'What the gap puts at risk, or what it blocks', None),
    ('Priority', 10, 'High / Medium / Low', 'Priority'),
    ('Related findings', 16, 'Finding IDs on the Review Trail', None),
    ('Source', 26, 'Script and run, finding, or meeting and date where it was raised', None),
    ('Owner', 18, 'Role that closes it', None),
    ('Status', 12, 'Open / In progress / Closed / Accepted', 'GapStatus'),
    ('Next step', 34, 'Next action, or why it was accepted', None),
    ('Raised', 12, 'Date raised', None),
]


def font(**kw):
    return Font(name=FONT, **kw)


def parse_register(path):
    rows = []
    with open(path, encoding='utf-8') as f:
        for line in f:
            if not re.match(r'^\| F-\d+ \|', line):
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if len(cells) < 14:
                continue
            fid, finding, source, count, pop, _result, sev, hitrust, zt, rec = cells[:10]
            sev = re.sub(r'\s*\(.*\)', '', sev).strip()
            if sev not in ('Critical', 'High', 'Medium', 'Low'):
                sev = ''
            rows.append([fid, finding, source.replace('`', ''), count, pop, sev, hitrust, zt, rec,
                         'findings\\%s.md' % fid, ''])
    return rows


def parse_gaps(path):
    rows = []
    with open(path, encoding='utf-8') as f:
        for line in f:
            if not re.match(r'^\| G-\d+ \|', line):
                continue
            cells = [x.strip().replace('`', '') for x in line.strip().strip('|').split('|')]
            if len(cells) < len(GAP_COLS):
                continue
            rows.append(cells[:len(GAP_COLS)])
    return rows


def build_gaps(wb, gaps, list_ref):
    gs = wb.create_sheet('Gaps', 1)
    hdr, light = '833C0B', 'FBE5D6'
    gs.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(GAP_COLS))
    h = gs.cell(row=1, column=1, value='Gap register  (control, capability, evidence, coverage and decision gaps; findings stay on the Review Trail)')
    h.font = font(bold=True, color='FFFFFF', size=11)
    h.fill = PatternFill('solid', fgColor=hdr)
    h.alignment = Alignment(horizontal='left', vertical='center')
    for ci, (name, width, comment, _l) in enumerate(GAP_COLS, start=1):
        cell = gs.cell(row=2, column=ci, value=name)
        cell.font = font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=hdr)
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BORDER
        cell.comment = Comment(comment, 'Gap register')
        gs.column_dimensions[get_column_letter(ci)].width = width
    gs.row_dimensions[1].height = 22
    gs.row_dimensions[2].height = 30
    first, last = 3, 3 + len(gaps) + 40
    for r in range(first, last + 1):
        data = gaps[r - first] if r - first < len(gaps) else None
        for ci in range(1, len(GAP_COLS) + 1):
            cell = gs.cell(row=r, column=ci, value=(data[ci - 1] or None) if data else None)
            cell.font = font(size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = BORDER
            cell.fill = PatternFill('solid', fgColor=light)
    for ci, (name, _w, _c, lst) in enumerate(GAP_COLS, start=1):
        col = get_column_letter(ci)
        if lst:
            v = DataValidation(type='list', formula1=list_ref(lst), allow_blank=True)
            v.error = 'Pick a value from the list'
            v.errorTitle = 'Gap register'
            gs.add_data_validation(v)
            v.add('%s%d:%s%d' % (col, first, col, last))
    col = lambda name: get_column_letter([g[0] for g in GAP_COLS].index(name) + 1)
    for s, colr in {'High': 'FF6600', 'Medium': 'FFC000', 'Low': '70AD47'}.items():
        gs.conditional_formatting.add('%s%d:%s%d' % (col('Priority'), first, col('Priority'), last), FormulaRule(
            formula=['%s%d="%s"' % (col('Priority'), first, s)], fill=PatternFill('solid', fgColor=colr),
            font=Font(name=FONT, bold=True, color='FFFFFF')))
    for s, colr in {'Closed': 'C6EFCE', 'Accepted': 'EDEDED', 'In progress': 'FFF2CC'}.items():
        gs.conditional_formatting.add('%s%d:%s%d' % (col('Status'), first, col('Status'), last), FormulaRule(
            formula=['%s%d="%s"' % (col('Status'), first, s)], fill=PatternFill('solid', fgColor=colr)))
    gs.freeze_panes = gs.cell(row=3, column=4)
    gs.auto_filter.ref = 'A2:%s%d' % (get_column_letter(len(GAP_COLS)), last)
    gs.sheet_view.zoomScale = 90
    gs.page_setup.orientation = 'landscape'
    gs.page_setup.fitToWidth = 1
    gs.page_setup.fitToHeight = 0
    gs.sheet_properties.pageSetUpPr.fitToPage = True
    gs.print_title_rows = '1:2'
    return col, first, last


def build(rows, out, title, as_of, gaps=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Review Trail'

    # ---- Lists (hidden) ----
    ls = wb.create_sheet('Lists')
    for ci, (name, values) in enumerate(LISTS.items(), start=1):
        ls.cell(row=1, column=ci, value=name).font = font(bold=True)
        for ri, v in enumerate(values, start=2):
            ls.cell(row=ri, column=ci, value=v).font = font()
    ls.sheet_state = 'hidden'

    def list_ref(name):
        ci = list(LISTS).index(name) + 1
        col = get_column_letter(ci)
        return "=Lists!$%s$2:$%s$%d" % (col, col, len(LISTS[name]) + 1)

    # ---- Review Trail headers ----
    col = 1
    colmap = {}
    for sec_title, owner, hdr, light, cols in SECTIONS:
        start = col
        for name, width, comment in cols:
            c = ws.cell(row=2, column=col, value=name)
            c.font = font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor=hdr)
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            c.border = BORDER
            c.comment = Comment(comment, 'Review trail')
            ws.column_dimensions[get_column_letter(col)].width = width
            colmap[name] = (col, light)
            col += 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
        h = ws.cell(row=1, column=start, value='%s  (%s)' % (sec_title, owner))
        h.font = font(bold=True, color='FFFFFF', size=11)
        h.fill = PatternFill('solid', fgColor=hdr)
        h.alignment = Alignment(horizontal='center', vertical='center')
    last_col = col - 1
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 42

    n = max(len(rows), 0)
    first, last = 3, max(3 + n - 1, 3) + 60  # 60 spare rows for new findings
    L = lambda name: get_column_letter(colmap[name][0])

    engineer_cols = [c[0] for c in SECTIONS[0][4]]
    for r in range(first, last + 1):
        data = rows[r - first] if r - first < n else None
        for name, (ci, light) in colmap.items():
            cell = ws.cell(row=r, column=ci)
            cell.font = font(size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = BORDER
            cell.fill = PatternFill('solid', fgColor=light)
        if data:
            for i, name in enumerate(engineer_cols):
                ws.cell(row=r, column=colmap[name][0], value=data[i] if data[i] != '' else None)
        # calculated columns
        idc, fnd = L('ID'), L('Finding')
        asev, align, dec = L('Architect severity'), L('Alignment'), L('Client decision')
        ws.cell(row=r, column=colmap['Stage'][0], value=(
            '=IF(%s%d="","",IF(%s%d<>"","Closed",IF(%s%d<>"","With client",IF(%s%d<>"","Architect reviewed","Draft"))))'
            % (idc, r, dec, r, align, r, asev, r)))
        ws.cell(row=r, column=colmap['Waiting on'][0], value=(
            '=IF(%s%d="","",IF(%s%d<>"","None",IF(%s%d<>"","Client",IF(%s%d<>"","Zero Trust Architect","Identity Architect"))))'
            % (idc, r, dec, r, align, r, asev, r)))
        ws.cell(row=r, column=colmap['Disposition'][0], value=(
            '=IF(%s%d="Accept and remediate","Remediation backlog (Jira story)",IF(%s%d="Accept risk","Exception register (owner, expiry, compensating control)",'
            'IF(%s%d="Defer","Status register (owner, risk, follow-on, target date)",IF(%s%d="Reject","Record the rationale; close",""))))'
            % (dec, r, dec, r, dec, r, dec, r)))
        for name in ('Stage', 'Waiting on', 'Disposition'):
            ws.cell(row=r, column=colmap[name][0]).font = font(size=10, bold=(name == 'Stage'))

    # ---- Data validation ----
    def dv(listname, colname):
        v = DataValidation(type='list', formula1=list_ref(listname), allow_blank=True)
        v.error = 'Pick a value from the list'
        v.errorTitle = 'Review trail'
        ws.add_data_validation(v)
        v.add('%s%d:%s%d' % (L(colname), first, L(colname), last))
    dv('Severity', 'Observed severity')
    dv('Severity', 'Architect severity')
    dv('Effort', 'Effort')
    dv('Gate', 'Target gate / sprint')
    dv('Pillar', 'ZT pillar')
    dv('Principle', 'ZT principle')
    dv('Maturity', 'Current maturity')
    dv('Maturity', 'Target maturity')
    dv('Alignment', 'Alignment')
    dv('Decision', 'Client decision')
    for dcol in ('Engineer date', 'Architect date', 'ZT date', 'Decision date'):
        v = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True)
        v.error = 'Enter a date'
        ws.add_data_validation(v)
        v.add('%s%d:%s%d' % (L(dcol), first, L(dcol), last))
        for r in range(first, last + 1):
            ws['%s%d' % (L(dcol), r)].number_format = 'yyyy-mm-dd'

    # ---- Conditional formatting on severity and stage ----
    sev_colors = {'Critical': 'C00000', 'High': 'FF6600', 'Medium': 'FFC000', 'Low': '70AD47'}
    for scol in ('Observed severity', 'Architect severity'):
        rng = '%s%d:%s%d' % (L(scol), first, L(scol), last)
        for s, colr in sev_colors.items():
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=['%s%d="%s"' % (L(scol), first, s)],
                fill=PatternFill('solid', fgColor=colr), font=Font(name=FONT, bold=True, color='FFFFFF')))
    rng = '%s%d:%s%d' % (L('Stage'), first, L('Stage'), last)
    for s, colr in {'Closed': 'C6EFCE', 'With client': 'FCE4D6', 'Architect reviewed': 'E4DFEC', 'Draft': 'DDEBF7'}.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s%d="%s"' % (L('Stage'), first, s)],
                                                       fill=PatternFill('solid', fgColor=colr)))

    ws.freeze_panes = ws.cell(row=3, column=3)
    ws.auto_filter.ref = 'A2:%s%d' % (get_column_letter(last_col), last)
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:2'

    # ---- Gaps ----
    gcol, gfirst, glast = build_gaps(wb, gaps or [], list_ref)

    # ---- Summary ----
    sm = wb.create_sheet('Summary', 0)
    sm.column_dimensions['A'].width = 34
    for c in 'BCDE':
        sm.column_dimensions[c].width = 14
    sm['A1'] = title
    sm['A1'].font = font(bold=True, size=14)
    sm['A2'] = 'Counts update as the Review Trail is filled in. As of: %s' % as_of
    sm['A2'].font = font(italic=True, size=9, color='595959')
    rt = "'Review Trail'!"

    def block(r0, heading, label_col_name, values):
        sm.cell(row=r0, column=1, value=heading).font = font(bold=True, color='FFFFFF')
        sm.cell(row=r0, column=1).fill = PatternFill('solid', fgColor='1F4E79')
        sm.cell(row=r0, column=2, value='Findings').font = font(bold=True, color='FFFFFF')
        sm.cell(row=r0, column=2).fill = PatternFill('solid', fgColor='1F4E79')
        rng = '%s$%d:$%s$%d' % (rt + '$' + L(label_col_name), first, L(label_col_name), last)
        r = r0 + 1
        for v in values:
            sm.cell(row=r, column=1, value=v).font = font()
            sm.cell(row=r, column=2, value='=COUNTIF(%s,A%d)' % (rng, r)).font = font()
            for cc in (1, 2):
                sm.cell(row=r, column=cc).border = BORDER
            r += 1
        sm.cell(row=r, column=1, value='Total').font = font(bold=True)
        sm.cell(row=r, column=2, value='=SUM(B%d:B%d)' % (r0 + 1, r - 1)).font = font(bold=True)
        return r + 2

    r = 4
    r = block(r, 'Stage', 'Stage', ['Draft', 'Architect reviewed', 'With client', 'Closed'])
    r = block(r, 'Observed severity', 'Observed severity', LISTS['Severity'])
    r = block(r, 'Architect severity', 'Architect severity', LISTS['Severity'])
    r = block(r, 'Client decision', 'Client decision', LISTS['Decision'])
    r = block(r, 'Zero Trust pillar', 'ZT pillar', LISTS['Pillar'])

    def gap_block(r0, heading, colname, values):
        for cc, txt in ((1, heading), (2, 'Gaps')):
            sm.cell(row=r0, column=cc, value=txt).font = font(bold=True, color='FFFFFF')
            sm.cell(row=r0, column=cc).fill = PatternFill('solid', fgColor='833C0B')
        rng = "Gaps!$%s$%d:$%s$%d" % (gcol(colname), gfirst, gcol(colname), glast)
        rr = r0 + 1
        for v in values:
            sm.cell(row=rr, column=1, value=v).font = font()
            sm.cell(row=rr, column=2, value='=COUNTIF(%s,A%d)' % (rng, rr)).font = font()
            for cc in (1, 2):
                sm.cell(row=rr, column=cc).border = BORDER
            rr += 1
        sm.cell(row=rr, column=1, value='Total').font = font(bold=True)
        sm.cell(row=rr, column=2, value='=SUM(B%d:B%d)' % (r0 + 1, rr - 1)).font = font(bold=True)
        return rr + 2
    r = gap_block(r, 'Gap type', 'Type', LISTS['GapType'])
    r = gap_block(r, 'Gap status', 'Status', LISTS['GapStatus'])

    # ---- Instructions ----
    ins = wb.create_sheet('Instructions', 0)
    ins.column_dimensions['A'].width = 36
    ins.column_dimensions['B'].width = 100
    lines = [
        ('Findings review trail', None),
        ('Purpose', 'One row per finding. Each role completes its own column group in order; the Stage, Waiting on and Disposition columns update on their own. A finding counts as decided only when its Stage is Closed.'),
        ('Order', '1. Identity Engineer drafts the finding with evidence (blue).  2. Identity Architect assesses and proposes the solution (green).  3. Zero Trust Architect confirms alignment and maturity (purple).  4. Client decision authority decides (orange).'),
        ('Which cells to edit', 'Only the shaded cells in your own column group on the Review Trail sheet. Columns with a dropdown accept list values only. Hover over a column header for what goes in it. Do not edit the grey Disposition formulas, except Tracking ID.'),
        ('Stage logic', 'Draft = Engineer section only. Architect reviewed = Architect severity set. With client = Alignment set. Closed = Client decision set.'),
        ('After the decision', 'Accept and remediate -> Jira story on the remediation backlog. Accept risk -> exception register (owner, expiry, compensating control). Defer -> status register (owner, risk, disposition, follow-on, target date). Reject -> record the rationale. Put the resulting ID in Tracking ID.'),
        ('Evidence', 'Each finding links to its evidence note (findings\\F-xx.md) in the approved evidence location: run folder, file SHA-256, run account, data-as-of time and the closure test. Never paste names or account details into this workbook.'),
        ('Gaps sheet', 'Everything that is a gap but not a counted finding: Control (a theme that groups findings), Capability (a process or tool that is missing, often raised in a meeting), Evidence (could not be assessed yet: access, inputs), Coverage (nothing measures it) and Decision (waiting on a decision). Link gaps to findings by ID; do not copy findings onto it. Close an Evidence gap when the run lands; it may then produce findings.'),
        ('Definitions', 'Severity scale, Zero Trust maturity levels, pillars, decisions and gap types: see the Definitions sheet.'),
        ('Example row (illustrative)', 'ID F-EX | Finding: security groups with no accountable owner | Source: 05 ad-groups.csv, OwnerStatus not OK | Count 42 of 310 | Observed severity Low | Architect severity Low; root cause: no owner field in the group request form; solution: make owner mandatory in the request form and assign owners for existing groups; effort M; gate G3 | ZT pillar Governance, Use least privilege, Initial -> Advanced, Aligned | Client decision Accept and remediate -> Remediation backlog (Jira story).'),
    ]
    for i, (k, v) in enumerate(lines, start=1):
        a = ins.cell(row=i, column=1, value=k)
        a.font = font(bold=True, size=14 if v is None else 10)
        a.alignment = Alignment(vertical='top')
        if v:
            b = ins.cell(row=i, column=2, value=v)
            b.font = font(size=10)
            b.alignment = Alignment(wrap_text=True, vertical='top')
    leg_r = len(lines) + 2
    ins.cell(row=leg_r, column=1, value='Colour key').font = font(bold=True)
    for j, (sec_title, owner, hdr, light, _cols) in enumerate(SECTIONS):
        a = ins.cell(row=leg_r + 1 + j, column=1, value=sec_title)
        a.fill = PatternFill('solid', fgColor=hdr)
        a.font = font(bold=True, color='FFFFFF', size=10)
        b = ins.cell(row=leg_r + 1 + j, column=2, value='Completed by: %s' % owner)
        b.fill = PatternFill('solid', fgColor=light)
        b.font = font(size=10)

    # ---- Definitions ----
    df = wb.create_sheet('Definitions')
    df.column_dimensions['A'].width = 26
    df.column_dimensions['B'].width = 90
    defs = [
        ('Severity', None),
        ('Critical', 'Reaches Tier 0 (domain or tenant control), or leaves a leaver or an admin with no second factor. Fix or formally accept within the current sprint.'),
        ('High', 'Privileged access held outside the intended model, or a gap that affects many accounts. Plan in the next sprint.'),
        ('Medium', 'Hygiene gap with limited reach, or one that another control partly covers. Schedule in the remediation backlog.'),
        ('Low', 'Housekeeping: naming, ownership records, clean-up of unused objects. Batch with routine work.'),
        ('Zero Trust maturity (CISA ZTMM v2.0)', None),
        ('Traditional', 'Manual, static configuration; siloed enforcement; limited visibility.'),
        ('Initial', 'Starting automation of attribute assignment and response; some cross-pillar integration.'),
        ('Advanced', 'Automated controls and centralized visibility across pillars; policies based on cross-pillar inputs.'),
        ('Optimal', 'Fully automated, just-in-time lifecycles and dynamic policies with continuous monitoring.'),
        ('Zero Trust principles', None),
        ('Verify explicitly', 'Authenticate and authorize on every request using all available signals.'),
        ('Use least privilege', 'Just-in-time and just-enough access; risk-based adaptive policies.'),
        ('Assume breach', 'Minimize blast radius, segment access, verify end to end, monitor.'),
        ('Client decision', None),
        ('Accept and remediate', 'The client accepts the finding and the proposed solution; it goes to the remediation backlog.'),
        ('Accept risk', 'The client accepts the residual risk without remediation; recorded in the exception register with owner, compensating control and expiry.'),
        ('Defer', 'Valid finding, not addressed in this phase; recorded in the status register with owner, risk, disposition, follow-on and target date. Never represented as implemented.'),
        ('Reject', 'The client does not accept the finding; the rationale is recorded.'),
        ('Gap type', None),
        ('Control', 'The environment falls short of the target; groups the findings that show it.'),
        ('Capability', 'A process, tool or function that does not exist or is not used (for example no access review process, licensed feature unused).'),
        ('Evidence', 'Not assessed yet because access, an export or an input is missing. Never read as a clean result.'),
        ('Coverage', 'Asked for by the scope but not measured by any script; needs another source.'),
        ('Decision', 'Work is waiting on a decision (platform direction, ownership, scope).'),
        ('Source', 'Severity scale: engagement register instructions. Maturity levels and pillars: CISA Zero Trust Maturity Model v2.0 (April 2023), https://www.cisa.gov/zero-trust-maturity-model'),  # guardrail:allow public CISA source URL, not an internal domain
    ]
    for i, (k, v) in enumerate(defs, start=1):
        a = df.cell(row=i, column=1, value=k)
        if v is None:
            a.font = font(bold=True, color='FFFFFF')
            a.fill = PatternFill('solid', fgColor='1F4E79')
            df.cell(row=i, column=2).fill = PatternFill('solid', fgColor='1F4E79')
        else:
            a.font = font(bold=True, size=10)
            b = df.cell(row=i, column=2, value=v)
            b.font = font(size=10)
            b.alignment = Alignment(wrap_text=True, vertical='top')

    wb.move_sheet('Lists', offset=10)
    wb.active = 0
    wb.save(out)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--register', default='')
    ap.add_argument('--gaps', default='')
    ap.add_argument('--out', required=True)
    ap.add_argument('--title', default='Findings review trail')
    ap.add_argument('--as-of', default='')
    a = ap.parse_args()
    rows = parse_register(a.register) if a.register else []
    gaps = parse_gaps(a.gaps) if a.gaps else []
    build(rows, a.out, a.title, a.as_of or 'template', gaps)
    print('rows', len(rows), 'gaps', len(gaps), '->', a.out)
